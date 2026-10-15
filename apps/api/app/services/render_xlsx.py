from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


PROVENANCE_NOTE = {
    "founder_stated": "Founder-stated",
    "benchmark": "Benchmark (not company fact)",
    "placeholder": "PLACEHOLDER — edit before investor send",
    "verified_metric": "Verified metric",
    "inferred": "Inferred from transcript",
}


def render_xlsx(
    *,
    path: Path,
    project_name: str,
    assumptions: list[dict],
    watermarked: bool,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"

    header = Font(bold=True, color="000000")
    thin = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    ws["A1"] = "PitchOS Financial Model"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = project_name
    ws["A3"] = "WATERMARK: Contains unverified assumptions" if watermarked else "Strict-ready export"
    ws["A3"].font = Font(bold=True, italic=True)

    headers = ["Key", "Label", "Value", "Unit", "Provenance", "Notes"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(5, col, h)
        cell.font = header
        cell.border = thin

    key_rows: dict[str, int] = {}
    for i, a in enumerate(assumptions):
        row = 6 + i
        key_rows[a["key"]] = row
        values = [
            a.get("key"),
            a.get("label"),
            float(a.get("value") or 0),
            a.get("unit"),
            a.get("provenance"),
            a.get("notes"),
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row, col, v)
            cell.border = thin
            if col == 5:
                prov = str(a.get("provenance") or "")
                cell.comment = Comment(PROVENANCE_NOTE.get(prov, prov), "PitchOS")
                # B&W friendly: bold placeholders
                if prov == "placeholder":
                    cell.font = Font(bold=True)

    # Model sheet
    ms = wb.create_sheet("Simple Model")
    ms["A1"] = "Simple 12-month projection (editable drivers on Assumptions)"
    ms["A1"].font = Font(bold=True, size=12)

    # Map common keys
    def ref(key: str, default: float = 0.0) -> str:
        if key in key_rows:
            return f"Assumptions!C{key_rows[key]}"
        return str(default)

    ms["A3"] = "Starting customers"
    ms["B3"] = f"={ref('starting_customers', 0)}"
    ms["A4"] = "ARPU (monthly)"
    ms["B4"] = f"={ref('arpu_monthly', 49)}"
    ms["A5"] = "Monthly burn"
    ms["B5"] = f"={ref('monthly_burn', 15000)}"
    ms["A6"] = "Gross margin"
    ms["B6"] = f"={ref('gross_margin', 0.75)}"
    ms["A7"] = "Raise amount"
    ms["B7"] = f"={ref('raise_amount', 0)}"
    ms["A8"] = "Implied runway (months, simplistic)"
    ms["B8"] = "=IF(B5=0,0,B7/B5)"

    ms["A10"] = "Month"
    ms["B10"] = "Customers (flat)"
    ms["C10"] = "MRR"
    ms["D10"] = "Gross profit"
    ms["E10"] = "Net (MRR*GM - burn)"
    for col in range(1, 6):
        ms.cell(10, col).font = header

    for m in range(1, 13):
        row = 10 + m
        ms.cell(row, 1, m)
        ms.cell(row, 2, "=B$3")  # flat for honesty — growth must be founder-stated later
        ms.cell(row, 3, f"=B{row}*B$4")
        ms.cell(row, 4, f"=C{row}*B$6")
        ms.cell(row, 5, f"=D{row}-B$5")

    ms["A25"] = "Note: Growth is intentionally flat unless founder confirms growth assumptions."
    ms["A26"] = "Provenance lives on the Assumptions sheet. Placeholders are bolded."

    # Widths
    for sheet in (ws, ms):
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 28
        sheet.column_dimensions["C"].width = 14
        sheet.column_dimensions["D"].width = 14
        sheet.column_dimensions["E"].width = 18
        sheet.column_dimensions["F"].width = 40

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
