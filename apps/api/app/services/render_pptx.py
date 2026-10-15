from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.ontology.claims import CLAIM_TYPE_LABELS, ClaimType, DECK_CLAIM_ORDER


SLIDE_PLAN: list[tuple[str, list[ClaimType]]] = [
    ("Title", []),
    ("Problem", [ClaimType.PROBLEM, ClaimType.WHY_NOW]),
    ("Insight", [ClaimType.INSIGHT, ClaimType.ICP]),
    ("Solution", [ClaimType.SOLUTION]),
    ("Market", [ClaimType.MARKET]),
    ("Product / Differentiation", [ClaimType.DIFFERENTIATION, ClaimType.COMPETITION]),
    ("Traction", [ClaimType.TRACTION]),
    ("Go-to-Market & Pricing", [ClaimType.GTM, ClaimType.PRICING, ClaimType.UNIT_ECONOMICS]),
    ("Team", [ClaimType.TEAM]),
    ("Ask & Use of Funds", [ClaimType.ASK, ClaimType.USE_OF_FUNDS, ClaimType.MILESTONES, ClaimType.RISK]),
]


def build_outline(project_name: str, claims: list[dict]) -> list[dict]:
    by_type = {c["claim_type"]: c for c in claims if c.get("status") != "rejected"}
    outline = []
    for title, types in SLIDE_PLAN:
        bullets = []
        for t in types:
            node = by_type.get(t.value)
            if node and (node.get("body") or "").strip():
                bullets.append(
                    {
                        "claim_type": t.value,
                        "label": CLAIM_TYPE_LABELS[t],
                        "text": node["body"],
                        "status": node.get("status"),
                        "provenance": node.get("provenance"),
                    }
                )
        outline.append({"title": title if title != "Title" else project_name, "bullets": bullets})
    return outline


def render_pptx(
    *,
    path: Path,
    project_name: str,
    founder_name: str,
    claims: list[dict],
    watermarked: bool,
    readiness_score: float,
) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    by_type = {c["claim_type"]: c for c in claims if c.get("status") != "rejected"}

    # Title
    _add_title_slide(
        prs,
        project_name,
        founder_name or "Founder",
        f"Fundraising readiness pack · score {readiness_score:.0f}/100"
        + (" · UNVERIFIED ASSUMPTIONS" if watermarked else ""),
    )

    for title, types in SLIDE_PLAN[1:]:
        bullets: list[str] = []
        for t in types:
            node = by_type.get(t.value)
            if not node:
                continue
            label = CLAIM_TYPE_LABELS[t]
            body = (node.get("body") or "").strip()
            if not body:
                continue
            tag = node.get("provenance") or ""
            status = node.get("status") or ""
            prefix = f"[{status}/{tag}] " if watermarked else ""
            bullets.append(f"{label}: {prefix}{body}")
        if not bullets:
            bullets = ["(No confirmed content for this slide yet)"]
        _add_bullets_slide(prs, title, bullets, watermarked=watermarked)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _set_run_font(run, size: int = 20, bold: bool = False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Calibri"


def _add_title_slide(prs: Presentation, title: str, subtitle: str, footer: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    _set_run_font(run, 40, True)

    box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(1))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    run2 = p2.add_run()
    run2.text = subtitle
    _set_run_font(run2, 22, False)

    box3 = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5))
    tf3 = box3.text_frame
    p3 = tf3.paragraphs[0]
    run3 = p3.add_run()
    run3.text = footer
    _set_run_font(run3, 12, False)


def _add_bullets_slide(prs: Presentation, title: str, bullets: list[str], watermarked: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.8), Inches(0.8))
    tf = title_box.text_frame
    run = tf.paragraphs[0].add_run()
    run.text = title
    _set_run_font(run, 28, True)

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.8), Inches(5.2))
    tfb = body.text_frame
    tfb.word_wrap = True
    for i, b in enumerate(bullets[:8]):
        p = tfb.paragraphs[0] if i == 0 else tfb.add_paragraph()
        p.text = f"• {b}"
        p.level = 0
        for r in p.runs:
            _set_run_font(r, 16, False)
        if not p.runs:
            run = p.add_run()
            run.text = f"• {b}"
            _set_run_font(run, 16, False)

    if watermarked:
        wm = slide.shapes.add_textbox(Inches(0.7), Inches(6.8), Inches(11.8), Inches(0.4))
        r = wm.text_frame.paragraphs[0].add_run()
        r.text = "WATERMARK: Contains unverified or unconfirmed assumptions — not investor-final."
        _set_run_font(r, 10, True)
